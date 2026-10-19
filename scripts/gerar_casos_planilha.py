"""Gera casos de teste preenchendo a planilha original e recalculando no LibreOffice.

Uso: python3 scripts/gerar_casos_planilha.py caminho/PLANILHA_TR.xlsx
Requer openpyxl e LibreOffice (soffice) instalados.
O resultado vai para src/motor/casos-planilha.json e é usado por acustica.test.ts.
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

CASOS = [
    {
        "nome": "sala de aula",
        "volume": 8 * 6 * 3,
        "ar": "Ar Umidade relativa 50%",
        "superficies": [
            ("piso", 48, "PISO VINÍLICO"),
            ("forro", 48, "GESSO LISO"),
            ("paredes", 70, "ALVENARIA LISA OU PINTADA"),
            ("janelas", 14, "VIDRO COMUM"),
            ("porta", 1.8, "MADEIRA PINTADA OU ENVERNIZADA"),
        ],
        "ocupacao": (40, "Cadeira com encosto e assento contraplacado (UNIDADE)", "ADULTO SENTADO INCLUINDO CADEIRA (UNIDADE)"),
    },
    {
        "nome": "estúdio muito absorvente",
        "volume": 5 * 4 * 2.8,
        "ar": None,
        "superficies": [
            ("piso", 20, "CARPETE"),
            ("forro", 20, "AMBI CF5 FORRO"),
            ("paredes", 50.4, "SONEX 50/125"),
        ],
        "ocupacao": None,
    },
    {
        "nome": "auditório",
        "volume": 30 * 20 * 8,
        "ar": "Ar- umidade relativa 40%",
        "superficies": [
            ("piso", 600, "CONCRETO"),
            ("forro", 600, "FORRO DE GESSO ACARTONADO COM OU SEM PREENCHIMENTO DE LÃ"),
            ("paredes laterais", 480, "TIJOLO APARENTE"),
            ("fundo", 160, "Cortina de veludo, esticada a 10cm da parede"),
            ("palco", 160, "MADEIRA PINTADA OU ENVERNIZADA"),
        ],
        "ocupacao": (400, "POLTRONA EM TECIDO VAZIA (UNIDADE)", "POLTRONA EM TECIDO CHEIA (UNIDADE)"),
    },
]


def nome_na_planilha(ws_mat, nome):
    """Devolve o texto exato da célula (a planilha tem espaços sobrando em alguns nomes)."""
    for (valor,) in ws_mat.iter_rows(min_row=2, max_col=1, values_only=True):
        if valor and " ".join(str(valor).split()) == nome:
            return valor
    raise KeyError(nome)


def main(caminho: str) -> None:
    raiz = Path(__file__).resolve().parent.parent
    saida = []
    with tempfile.TemporaryDirectory() as tmp:
        for n, caso in enumerate(CASOS):
            wb = openpyxl.load_workbook(caminho)
            mat, ws = wb["MATERIAIS"], wb["CÁLCULO"]
            ws["E3"] = caso["volume"]
            ws["F3"] = nome_na_planilha(mat, caso["ar"]) if caso["ar"] else 0
            for i, (_, area, material) in enumerate(caso["superficies"]):
                ws[f"E{5 + i}"] = area
                ws[f"F{5 + i}"] = nome_na_planilha(mat, material)
            if caso["ocupacao"]:
                qtd, vazio, ocupado = caso["ocupacao"]
                ws["E34"] = qtd
                for linha, material in ((34, vazio), (35, ocupado), (36, vazio), (37, ocupado)):
                    ws[f"F{linha}"] = nome_na_planilha(mat, material)
            entrada = Path(tmp) / f"caso{n}.xlsx"
            wb.save(entrada)
            saida_dir = Path(tmp) / "out"
            subprocess.run(
                ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(saida_dir), str(entrada)],
                check=True, capture_output=True,
            )
            calc = openpyxl.load_workbook(saida_dir / entrada.name, data_only=True)["CÁLCULO"]
            linha = lambda r: [calc.cell(r, c).value for c in range(20, 26)]  # T..Y
            saida.append({
                **{k: caso[k] for k in ("nome", "volume", "ar")},
                "superficies": [{"nome": s, "quantidade": a, "material": m} for s, a, m in caso["superficies"]],
                "ocupacao": dict(zip(("quantidade", "vazio", "ocupado"), caso["ocupacao"])) if caso["ocupacao"] else None,
                "esperado": {
                    "absorcao": {"vazio": linha(2), "meio": linha(3), "cheio": linha(4)},
                    "alfaMedio": {"vazio": linha(5), "meio": linha(6), "cheio": linha(7)},
                    "sabine": {"vazio": linha(8), "meio": linha(9), "cheio": linha(10)},
                    "eyring": {"vazio": linha(11), "meio": linha(12), "cheio": linha(13)},
                    "relacaoAV": [calc.cell(5, c).value for c in range(28, 34)],  # AB..AG
                },
            })
    destino = raiz / "src" / "motor" / "casos-planilha.json"
    destino.write_text(json.dumps(saida, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{len(saida)} casos gravados em {destino}")


if __name__ == "__main__":
    main(sys.argv[1])
